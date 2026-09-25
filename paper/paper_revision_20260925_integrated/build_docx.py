"""Build both editable review documents from the single Markdown render manifest.

Run with the bundled documents-runtime Python. Equations are embedded images
whose editable formula source remains in Markdown; figures are embedded too.
"""
from pathlib import Path
import json,re
from PIL import Image
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

HERE=Path(__file__).resolve().parent
manifest=json.loads((HERE/'build/render-manifest.json').read_text(encoding='utf-8'))
WIDTH=6.82

def element(name,**attrs):
    e=OxmlElement('w:'+name)
    for k,v in attrs.items():e.set(qn('w:'+k),str(v))
    return e

def inline(p,text):
    pattern=r'(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*|`[^`]+`)'
    for part in re.split(pattern,text):
        link=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',part)
        if link:
            label,url=link.groups()
            if not re.match(r'https?://',url):url=(HERE/url).resolve().as_uri()
            h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
            r=OxmlElement('w:r');pr=OxmlElement('w:rPr');pr.append(element('color',val='176B87'));r.append(pr)
            t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)
        else:
            bold=part.startswith('**') and part.endswith('**')
            code=part.startswith('`') and part.endswith('`')
            r=p.add_run(part[2:-2] if bold else part[1:-1] if code else part)
            if bold:r.bold=True
            if code:r.font.name='Consolas';r.font.size=Pt(8)

def number(value):
    if re.fullmatch(r'[-+]?\d*\.?\d+(?:e[-+]?\d+)?',value,re.I):
        v=float(value)
        return str(int(v)) if v==int(v) and abs(v)<100000 else f'{v:.6g}'
    return value

def widths_for(rows):
    n=len(rows[0]);h=rows[0]
    if n==2:return [WIDTH*.31,WIDTH*.69]
    if h[0] in ('Statement','Source') and n==4:
        ratios=[1.6,1.15,1.35,1.6]
    elif h[:3]==['Reference','Seed','Role']:
        ratios=[.85,.62,.72]+[1.]*(n-3)
    elif h[:2]==['Seed','Role']:ratios=[.65,.75]+[1.]*(n-2)
    elif h[0] in ('Role','Method','Pool','pool'):ratios=[1.7]+[1.]*(n-1)
    elif h[0] in ('Arm','State','Comparison'):ratios=[1.1]+[1.]*(n-1)
    else:ratios=[1.]*n
    return [WIDTH*r/sum(ratios) for r in ratios]

def add_table(doc,rows):
    n=len(rows[0]);widths=widths_for(rows);tab=doc.add_table(rows=0,cols=n)
    tab.alignment=WD_TABLE_ALIGNMENT.CENTER;tab.autofit=False
    for col,width in zip(tab.columns,widths):col.width=Inches(width)
    pr=tab._tbl.tblPr;borders=OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        borders.append(element(side,val='single',sz='4',color='D9D9D9'))
    pr.append(borders)
    for j,values in enumerate(rows):
        row=tab.add_row();rp=row._tr.get_or_add_trPr();rp.append(element('cantSplit'))
        if j==0:rp.append(element('tblHeader'))
        for k,(cell,value) in enumerate(zip(row.cells,values)):
            cell.width=Inches(widths[k]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
            for side,val in [('top',75),('bottom',75),('left',65),('right',65)]:m.append(element(side,w=val,type='dxa'))
            cp.append(m);cp.append(element('shd',fill='182A3A' if j==0 else 'F0F4F7' if j%2 else 'FFFFFF'))
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0
            p.paragraph_format.keep_with_next=(j==0)
            raw=number(value) if j else value.replace('_',' ')
            inline(p,raw)
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if (k==0 or len(raw)>35) else WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name='Arial';r.font.size=Pt(7.3 if n>=9 else 7.8)
                if j==0:r.bold=True;r.font.color.rgb=RGBColor.from_string('FFFFFF')
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

summary={}
for name,blocks in manifest['documents'].items():
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929)
    sec.left_margin=sec.right_margin=Inches((8.2677-WIDTH)/2)
    sec.top_margin=Inches(.70);sec.bottom_margin=Inches(.72)
    sec.header_distance=Inches(.31);sec.footer_distance=Inches(.31)
    for style in ('Normal','Title','Heading 1','Heading 2','Heading 3','Caption'):
        st=doc.styles[style];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        st.font.size=Pt(10.5);st.paragraph_format.widow_control=True
        for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
        for fonts in st.element.iter(qn('w:rFonts')):
            for key in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
                fonts.attrib.pop(qn('w:'+key),None)
            fonts.set(qn('w:ascii'),'Times New Roman');fonts.set(qn('w:hAnsi'),'Times New Roman')
    normal=doc.styles['Normal'];normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.08
    for st,size in [('Title',19),('Heading 1',13),('Heading 2',11.5),('Heading 3',10.8)]:
        s=doc.styles[st];s.font.size=Pt(size);s.font.bold=True
        s.paragraph_format.keep_with_next=True;s.paragraph_format.space_before=Pt(12);s.paragraph_format.space_after=Pt(6)
    doc.styles['Caption'].font.size=Pt(8.3);doc.styles['Caption'].font.italic=False
    hp=sec.header.paragraphs[0];hp.text='PINN electrothermal phase reconstruction';hp.runs[0].font.size=Pt(8)
    fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=fp.add_run('Integrated review 25 September 2026  |  ');r.font.size=Pt(8)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');fp._p.append(fld)
    for i,b in enumerate(blocks):
        kind=b['type']
        if kind=='heading':
            p=doc.add_paragraph(style='Title' if b['level']==1 else 'Heading '+str(min(b['level']-1,3)))
            inline(p,b['text'])
        elif kind=='paragraph':
            text=b['text'];caption=text.startswith(('Figure ','Table '))
            p=doc.add_paragraph(style='Caption' if caption else 'Normal');inline(p,text)
            if text.startswith('Table '):p.paragraph_format.keep_with_next=True
            if re.match(r'^\[\d+\]',text):
                for r in p.runs:r.font.size=Pt(8.5)
                p.paragraph_format.space_after=Pt(4)
        elif kind=='table':add_table(doc,b['rows'])
        elif kind in ('image','equation'):
            with Image.open(b['path']) as im:w,h=im.size
            width=min(WIDTH,w/220) if kind=='equation' else min(WIDTH,7.8*w/h)
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3)
            p.add_run().add_picture(b['path'],width=Inches(width))
            if kind=='image' and i+1<len(blocks) and blocks[i+1].get('text','').startswith('Figure '):p.paragraph_format.keep_with_next=True
    doc.core_properties.title=blocks[0]['text'];doc.core_properties.author=''
    doc.core_properties.subject='Unified manuscript from frozen evidence and a bounded conditional diagnostic'
    path=HERE/f'{name}.docx';doc.save(path)
    summary[name]=dict(output=str(path),tables=len(doc.tables),images=len(doc.inline_shapes),paragraphs=len(doc.paragraphs))
(HERE/'build/docx-build-summary.json').write_text(json.dumps(summary,indent=2),encoding='utf-8')
print(json.dumps(summary,indent=2))
